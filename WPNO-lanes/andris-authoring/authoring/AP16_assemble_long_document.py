#!/usr/bin/env python3
"""
AP-16 — Sadaļu montāža garam dokumentam (section assembly for long documents).

Takes the WPNO letterhead master (v6 .dotx) plus a structured JSON outline
(Rubrum, Antraege, Sections, Anlagen-Definitionen) and assembles ONE long,
internally-consistent document:
  - Rubrum block
  - Numbered Antrag block (real Word auto-numbering)
  - Gliederung headings (Gliederung1/Gliederung2 styles, which already carry
    outlineLvl 0/1 -> a native Word TOC field picks them up automatically)
  - Section bodies in Fliesstext, each carrying its own Anlage references
  - An automatically-collected Anlagenverzeichnis at the end (only Anlagen
    that are actually referenced in the body text are listed — this *is*
    the cross-reference consistency check per Implementation Paper 22.2/22.4)

SYNTHETIC test data only, per Chapter 13 privacy discipline — no real
mandate data.

Usage:
    python3 AP16_assemble_long_document.py <template.dotx> <outline.json> <output.docx>
"""
import sys
import json
import re
import zipfile
import shutil
import os
import subprocess

def dotx_to_docx_bytes(dotx_path, tmp_docx_path):
    zin = zipfile.ZipFile(dotx_path, 'r')
    zout = zipfile.ZipFile(tmp_docx_path, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == '[Content_Types].xml':
            text = data.decode('utf-8')
            text = text.replace(
                'application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml',
                'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml'
            )
            data = text.encode('utf-8')
        zout.writestr(item, data)
    zout.close()

NUMBERING_XML = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:abstractNum w:abstractNumId="100">
<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/>
<w:lvlJc w:val="start"/><w:pPr><w:ind w:start="357" w:hanging="357"/></w:pPr></w:lvl>
</w:abstractNum>
<w:num w:numId="100"><w:abstractNumId w:val="100"/></w:num>
</w:numbering>'''

def add_numbering_part(docx_path):
    tmp = docx_path + '.tmp'
    zin = zipfile.ZipFile(docx_path, 'r')
    names = zin.namelist()
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == '[Content_Types].xml':
            text = data.decode('utf-8')
            if 'numbering.xml' not in text:
                text = text.replace(
                    '</Types>',
                    '<Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/></Types>'
                )
            data = text.encode('utf-8')
        if item.filename == 'word/_rels/document.xml.rels':
            text = data.decode('utf-8')
            if 'numbering.xml' not in text:
                text = text.replace(
                    '</Relationships>',
                    '<Relationship Id="rIdNumbering100" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering" Target="numbering.xml"/></Relationships>'
                )
            data = text.encode('utf-8')
        zout.writestr(item, data)
    if 'word/numbering.xml' not in names:
        zout.writestr('word/numbering.xml', NUMBERING_XML)
    zout.close()
    shutil.move(tmp, docx_path)

def apply_numpr(paragraph, num_id=100, ilvl=0):
    from docx.oxml.ns import qn
    pPr = paragraph._p.get_or_add_pPr()
    numPr = pPr.makeelement(qn('w:numPr'), {})
    ilvl_el = numPr.makeelement(qn('w:ilvl'), {qn('w:val'): str(ilvl)})
    numId_el = numPr.makeelement(qn('w:numId'), {qn('w:val'): str(num_id)})
    numPr.append(ilvl_el)
    numPr.append(numId_el)
    pPr.append(numPr)

def build_toc_paragraph(doc):
    """Insert a native Word TOC field. Word/LibreOffice will populate it
    from headings whose style carries an outlineLvl (Gliederung1/2 already do)."""
    from docx.oxml.ns import qn
    p = doc.add_paragraph()
    p.style = doc.styles['Gliederung1']
    run_heading = p.add_run("Inhaltsverzeichnis")
    p2 = doc.add_paragraph()
    r = p2.add_run()
    fldChar_begin = r._r.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'begin'})
    instrText = r._r.makeelement(qn('w:instrText'), {qn('xml:space'): 'preserve'})
    instrText.text = ' TOC \\o "1-2" \\h \\z \\u '
    fldChar_sep = r._r.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'separate'})
    fldChar_end = r._r.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'end'})
    r._r.append(fldChar_begin)
    r._r.append(instrText)
    r._r.append(fldChar_sep)
    r._r.append(fldChar_end)
    return p2

def main():
    if len(sys.argv) != 4:
        print("Usage: python3 AP16_assemble_long_document.py <template.dotx> <outline.json> <output.docx>")
        sys.exit(1)

    template_path, outline_path, output_path = sys.argv[1:4]

    with open(outline_path, encoding='utf-8') as f:
        outline = json.load(f)

    tmp_docx = output_path + '.building.docx'
    dotx_to_docx_bytes(template_path, tmp_docx)

    import docx
    d = docx.Document(tmp_docx)

    # Clear existing Nutzungshinweis instructional content
    for p in list(d.paragraphs):
        p._p.getparent().remove(p._p)

    # 1. Rubrum
    r = outline['rubrum']
    d.add_paragraph('Rubrum', style='RubrumFett')
    d.add_paragraph(r['gericht'], style='Fliesstext')
    d.add_paragraph(f"Az.: {r['aktenzeichen']}", style='Fliesstext')
    d.add_paragraph('In dem Rechtsstreit', style='Fliesstext')
    d.add_paragraph(r['klaeger'], style='Fliesstext')
    d.add_paragraph(f"Prozessbevollmaechtigte: {r['prozessbevollmaechtigte_klaeger']}", style='Fliesstext')
    d.add_paragraph('gegen', style='Fliesstext')
    d.add_paragraph(r['beklagte'], style='Fliesstext')
    d.add_paragraph(f"Prozessbevollmaechtigte: {r['prozessbevollmaechtigte_beklagte']}", style='Fliesstext')
    d.add_paragraph(r['streitgegenstand'], style='Fliesstext')

    # 2. Antrag (numbered)
    d.add_paragraph('Antrag', style='Antrag')
    d.add_paragraph('Namens und in Vollmacht der Klaegerin wird beantragt,', style='Fliesstext')
    antrag_paragraphs = []
    for a in outline['antraege']:
        p = d.add_paragraph(a, style='Antrag')
        antrag_paragraphs.append(p)

    # 3. Gliederung / TOC
    build_toc_paragraph(d)

    # 4. Sections
    style_by_level = {1: 'Gliederung1', 2: 'Gliederung2'}
    all_body_text = []
    for sec in outline['sections']:
        style_name = style_by_level.get(sec['level'], 'Gliederung2')
        d.add_paragraph(sec['heading'], style=style_name)
        for para_text in sec['paragraphs']:
            d.add_paragraph(para_text, style='Fliesstext')
            all_body_text.append(para_text)

    # 5. Cross-reference consistency: only list Anlagen that are actually
    #    referenced in the body text (this IS the consistency check).
    combined_text = " ".join(all_body_text)
    referenced = sorted(set(re.findall(r'Anlage\s+(K\d+)', combined_text)))
    defined = outline['anlagen_definitionen']

    missing_definitions = [k for k in referenced if k not in defined]
    unused_definitions = [k for k in defined if k not in referenced]

    d.add_paragraph('Anlagenverzeichnis', style='Anlagenverzeichnis')
    for k in referenced:
        desc = defined.get(k, '[FEHLENDE DEFINITION]')
        d.add_paragraph(f"Anlage {k} - {desc}", style='Anlage')

    d.save(tmp_docx)
    add_numbering_part(tmp_docx)

    # Re-open to apply numPr to Antrag paragraphs (must happen after numbering.xml part exists)
    d2 = docx.Document(tmp_docx)
    count = 0
    for p in d2.paragraphs:
        if p.style is not None and p.style.name == 'Antrag' and p.text.strip() and not p.text.strip() == 'Antrag':
            apply_numpr(p)
            count += 1
    d2.save(output_path)
    os.remove(tmp_docx)

    print(f"OK: assembled document written to {output_path}")
    print(f"    Sections: {len(outline['sections'])}, Antrag items numbered: {count}")
    print(f"    Anlagen referenced in text: {referenced}")
    if missing_definitions:
        print(f"    WARNUNG: referenzierte Anlagen ohne Definition: {missing_definitions}")
    if unused_definitions:
        print(f"    INFO: definierte, aber nicht referenzierte Anlagen (nicht aufgenommen): {unused_definitions}")
    print(f"    Konsistenzpruefung: {'PASS' if not missing_definitions else 'FAIL'}")

if __name__ == '__main__':
    main()
